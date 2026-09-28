#!/usr/bin/env python3
"""
Spreadsheet tools: `write_spreadsheet` creates a real Excel workbook (.xlsx) —
or a CSV — from rows of data.

Asked for "an Excel sheet of the tags", the model had no way to make one: it
tried pandas through the shell (not installed), then fell back to a CSV. The
obvious alternative was worse — write_file to tags.xlsx writes TEXT into a file
named .xlsx, which Excel refuses to open. openpyxl is already a dependency (for
reading workbooks), so writing one costs nothing new.

The schema asks for `rows` as TEXT - CSV lines or a markdown table. Measured on
gemma4:12b, one live run each: with an array-of-arrays schema the turn ended in
an empty reply straight after view_image; with a string schema the same prompt
produced a correct call first time, all four tags and their values exact.
Re-measure before switching it back. Rows are still accepted however a small
model sends them: CSV / TSV / markdown-table text, a JSON string, a list of
lists, a list of records (dicts), a dict of columns, or a flat list (one
column). `rows_from_any` normalises all of these, and the same parser lets
write_file turn text aimed at a .xlsx path into a real workbook instead of a
corrupt one.

Cells: numbers stay numbers; numeric-looking strings become numbers unless they
carry a leading zero (part numbers, "007"), so Excel can sum them; values
starting with "=" are Excel formulas, as they would be typed into Excel;
everything else is text. The header row is bold and frozen and columns are
sized to their content.

A workbook is always written fresh. openpyxl drops charts and images from
workbooks it re-saves, so an existing file is never edited in place: replacing
it needs overwrite=true, and the old file is backed up first.
"""

import csv
import io
import json
import os
import re
from itertools import zip_longest
from pathlib import Path
from typing import Any, Dict, List, Tuple

_NUMBER = re.compile(r"^-?(0|[1-9]\d*)(\.\d+)?$")
_MD_SEPARATOR = re.compile(r"^:?-{2,}:?$")
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
MAX_ROWS = 100_000
MAX_COLS = 500
_MIN_WIDTH, _MAX_WIDTH = 8, 60


# ---------------------------------------------------------------------------
# Getting rows out of whatever the model sent
# ---------------------------------------------------------------------------

def rows_from_any(value: Any) -> List[List[Any]]:
    """Normalise model input into a list of rows (lists of cell values)."""
    if value is None:
        return []
    if isinstance(value, str):
        return rows_from_text(value)
    if isinstance(value, dict):
        if value and all(isinstance(v, (list, tuple)) for v in value.values()):
            # A dict of columns: {"Tag": [...], "Value": [...]}
            headers = list(value.keys())
            return [headers] + [list(r) for r in zip_longest(*value.values(), fillvalue=None)]
        return [list(value.keys()), list(value.values())]      # a single record
    if isinstance(value, (list, tuple)):
        items = list(value)
        if not items:
            return []
        if all(isinstance(r, dict) for r in items):
            headers: List[Any] = []
            for r in items:
                for k in r:
                    if k not in headers:
                        headers.append(k)
            return [headers] + [[r.get(h) for h in headers] for r in items]
        # Rows as lists; a bare scalar becomes a one-cell row (a flat list is one column).
        return [list(r) if isinstance(r, (list, tuple)) else [r] for r in items]
    return [[value]]


def rows_from_text(text: str) -> List[List[Any]]:
    """JSON, a markdown table, or CSV/TSV text -> rows."""
    text = (text or "").strip()
    if not text:
        return []
    if text[:1] in "[{":
        try:
            parsed = json.loads(text)
        except json.JSONDecodeError:
            parsed = None
        if isinstance(parsed, (list, dict)):
            return rows_from_any(parsed)

    lines = [ln for ln in text.splitlines() if ln.strip()]
    if lines and all(ln.strip().startswith("|") for ln in lines):
        rows = []
        for ln in lines:
            cells = [c.strip() for c in ln.strip().strip("|").split("|")]
            if cells and all(_MD_SEPARATOR.match(c) for c in cells if c):
                continue                                          # |---|:---:|
            rows.append(cells)
        return rows

    sample = "\n".join(lines[:50])
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = "\t" if "\t" in sample else ","
    reader = csv.reader(io.StringIO("\n".join(lines)), delimiter=delimiter)
    return [[c.strip() for c in row] for row in reader if any(c.strip() for c in row)]


# ---------------------------------------------------------------------------
# Writing
# ---------------------------------------------------------------------------

def _cell(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    text = str(value)
    stripped = text.strip()
    if _NUMBER.match(stripped):
        try:
            return float(stripped) if "." in stripped else int(stripped)
        except ValueError:
            return text
    return text


def _sheet_name(name: Any) -> str:
    clean = _BAD_SHEET_CHARS.sub(" ", str(name or "").strip()).strip().strip("'")
    return (clean or "Sheet1")[:31]


def save_rows(path: Path, rows: List[List[Any]], sheet: Any = "Sheet1") -> Tuple[int, int]:
    """Write rows to a fresh .xlsx (or .csv). Returns (rows, columns) written."""
    rows = [list(r)[:MAX_COLS] for r in rows[:MAX_ROWS]]
    width = max((len(r) for r in rows), default=0)

    if path.suffix.lower() == ".csv":
        # utf-8-sig: the BOM is how Excel knows a CSV is UTF-8 - without it,
        # accented names and symbols open as mojibake.
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            for r in rows:
                writer.writerow(["" if c is None else c for c in r])
        return len(rows), width

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = _sheet_name(sheet)
    for r in rows:
        ws.append([_cell(c) for c in r])
    if rows:
        for c in ws[1]:
            c.font = Font(bold=True)
        if len(rows) > 1:
            ws.freeze_panes = "A2"
    widths: Dict[int, int] = {}
    for r in rows:
        for i, c in enumerate(r):
            widths[i] = max(widths.get(i, 0), len(str(c)) if c is not None else 0)
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i + 1)].width = min(_MAX_WIDTH, max(_MIN_WIDTH, w + 2))
    # openpyxl never computes formulas, so a formula cell carries no cached value;
    # it relies on openpyxl's default fullCalcOnLoad=1 for Excel and Calc to
    # recompute on open (verified with LibreOffice: =SUM exported as 4510).
    wb.save(str(path))
    return len(rows), width


# ---------------------------------------------------------------------------
# The tool
# ---------------------------------------------------------------------------

def write_spreadsheet(inp: Dict[str, Any]) -> str:
    """Create a real .xlsx (or .csv) from rows. Returns text, never raises."""
    from .file_tools import _backup_existing, _is_write_allowed, _write_denied_msg

    raw = str(inp.get("path", "")).strip().strip('"')
    if not raw:
        return "Error: 'path' is required"
    target = Path(os.path.expanduser(raw))
    ext = target.suffix.lower()
    if not ext:
        target, ext = target.with_suffix(".xlsx"), ".xlsx"      # "tags" -> "tags.xlsx"
    if ext not in (".xlsx", ".csv"):
        return (f"Error: write_spreadsheet writes Excel .xlsx (or .csv) files, not {ext}. "
                f"Use a path ending in .xlsx, e.g. {target.with_suffix('.xlsx').name}.")
    if not _is_write_allowed(target):
        return _write_denied_msg(target)

    source = inp.get("rows")
    if source is None:
        source = inp.get("data")          # tolerated: what small models sometimes call it
    try:
        rows = rows_from_any(source)
    except Exception as e:
        return f"Error: could not understand 'rows' ({e}). Pass a list of rows, each a list of cell values."
    headers = inp.get("headers") or inp.get("columns")
    if isinstance(headers, (list, tuple)) and headers and not any(isinstance(h, (list, tuple, dict)) for h in headers):
        if not rows or [str(c) for c in rows[0]] != [str(h) for h in headers]:
            rows = [list(headers)] + rows
    if not rows:
        return "Error: 'rows' is empty - nothing to write."

    existed = target.exists()
    if existed and not inp.get("overwrite"):
        return (f"Error: {target} already exists. Pass overwrite=true to replace it (the old file is "
                "backed up first), or choose a new file name.")
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        if existed:
            _backup_existing(target)
        n, width = save_rows(target, rows, inp.get("sheet") or "Sheet1")
    except PermissionError:
        return f"Error: could not write {target} - it is probably open in Excel. Close it and try again."
    except Exception as e:
        return f"Error writing spreadsheet: {e}"

    kind = "Excel workbook" if ext == ".xlsx" else "CSV file (opens in Excel)"
    first = ", ".join(str(c) for c in rows[0][:6]) + (", ..." if len(rows[0]) > 6 else "")
    text = f"Created {kind} {target} - {n} row(s) x {width} column(s). First row: {first}."
    if ext == ".xlsx" and any(isinstance(c, str) and c.startswith("=") for r in rows for c in r):
        text += " Formulas are calculated when the file is opened in Excel."
    if existed:
        text += " The previous file was backed up to .gemma/backups."
    return text
