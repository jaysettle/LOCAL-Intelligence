"""Tests for vision: view_image, pictures inside documents, and how images reach Ollama.

Fixtures are REAL files: a .docx / .pptx / .xlsx with a picture embedded by the
same libraries Office users' documents come from, an image-only PDF written by
Pillow (a stand-in for a scan), crafted OpenDocument and EPUB containers. The
model side is a fake Ollama that records each payload, so the tests pin what
actually goes on the wire.
"""
import base64
import copy
import io
import json
import zipfile

import pytest
from PIL import Image
from rich.console import Console

from gemma_cli import agent, config, review, sessions
from gemma_cli import main as main_mod
from gemma_cli.tools import doc_tools, file_tools, image_tools
from gemma_cli.tools.image_tools import (
    MAX_SIDE, TILE_ABOVE, ToolOutput, describe_embedded, encode_for_model, encode_user_images,
    list_embedded_images, page_ranges, view_image,
)


# --- fixtures ----------------------------------------------------------------

def _png(path, size=(200, 100), color=(200, 30, 30), mode="RGB"):
    Image.new(mode, size, color).save(path)
    return path


def _decode(b64):
    return Image.open(io.BytesIO(base64.b64decode(b64)))


@pytest.fixture
def folder(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    file_tools.set_allowed_write_roots([tmp_path])
    monkeypatch.setattr(config, "CONFIG_PATH", tmp_path / "nope.yaml")
    return tmp_path


def _docx(path, pictures, text="Dashboard tags for Canary"):
    import docx
    d = docx.Document()
    if text:
        d.add_paragraph(text)
    for p in pictures:
        d.add_picture(str(p))
    d.save(str(path))
    return path


def _pptx(path, picture):
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(picture), 0, 0)
    prs.save(str(path))
    return path


def _xlsx(path, picture):
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    wb = openpyxl.Workbook()
    wb.active["A1"] = "tags"
    wb.active.add_image(XLImage(str(picture)), "B2")
    wb.save(str(path))
    return path


def _scan_pdf(path, pages=1):
    frames = [Image.new("RGB", (400, 300), (255, 255, 255)) for _ in range(pages)]
    frames[0].save(str(path), save_all=True, append_images=frames[1:])
    return path


# --- encoding pixels ----------------------------------------------------------

def test_small_image_is_one_png(folder):
    enc = encode_for_model(_png(folder / "a.png"))
    assert len(enc.images) == 1 and not enc.tiled
    assert enc.size == (200, 100)
    assert _decode(enc.images[0]).format == "PNG"


def test_large_image_is_overview_plus_four_tiles(folder):
    enc = encode_for_model(_png(folder / "big.png", size=(TILE_ABOVE + 400, 900)))
    assert enc.tiled and len(enc.images) == 5
    assert "overview plus 4 zoomed tiles" in enc.describe()


def test_tiles_overlap_and_cover_the_image(folder):
    im = Image.new("RGB", (2000, 1000), "white")
    tiles = image_tools._tiles(im)
    assert len(tiles) == 4
    for t in tiles:
        assert t.size[0] > 1000 and t.size[1] > 500     # half plus overlap
    # A 150px-wide label straddling the vertical seam lies whole inside a tile.
    left_end = tiles[0].size[0]
    right_start = 2000 - tiles[1].size[0]
    assert left_end - right_start >= 150


def test_transparency_is_flattened_onto_white(folder):
    p = folder / "t.png"
    Image.new("RGBA", (50, 50), (0, 0, 0, 0)).save(p)
    px = _decode(encode_for_model(p).images[0]).convert("RGB").getpixel((10, 10))
    assert px == (255, 255, 255)


@pytest.mark.parametrize("fmt,ext", [("BMP", ".bmp"), ("GIF", ".gif"), ("TIFF", ".tif"), ("JPEG", ".jpg")])
def test_other_formats_become_png(folder, fmt, ext):
    p = folder / f"x{ext}"
    Image.new("RGB", (60, 40), (10, 200, 10)).save(p, fmt)
    assert _decode(encode_for_model(p).images[0]).format == "PNG"


def test_oversized_images_are_capped(folder):
    enc = encode_for_model(_png(folder / "huge.png", size=(5000, 3000)))
    for b64 in enc.images:
        assert max(_decode(b64).size) <= MAX_SIDE


def test_undecodable_bytes_raise_a_readable_error():
    with pytest.raises(ValueError, match="not an image format"):
        encode_for_model(b"this is not an image")


def test_user_attachments_notes(folder):
    imgs, note = encode_user_images([folder / "missing.png"])
    assert imgs == [] and "was not found" in note
    imgs, note = encode_user_images([_png(folder / "small.png")])
    assert len(imgs) == 1 and note == ""
    imgs, note = encode_user_images([_png(folder / "shot.png", size=(1920, 1080))])
    assert len(imgs) == 5 and "zoomed tiles" in note and "shot.png" in note
    bad = folder / "bad.png"
    bad.write_bytes(b"nope")
    imgs, note = encode_user_images([bad])
    assert imgs == [] and "no image is attached" in note


# --- pictures inside documents ----------------------------------------------

def test_docx_pictures_listed_in_natural_order(folder):
    pics = [_png(folder / f"p{i}.png", color=(i * 20, 0, 0)) for i in range(1, 12)]
    doc = _docx(folder / "d.docx", pics)
    found = list_embedded_images(doc, "docx")
    assert len(found) == 11
    assert found[1].label == "image2.png" and found[10].label == "image11.png"   # image11 after image2
    assert _decode(encode_for_model(found[0].load()).images[0]).size == (200, 100)


def test_pptx_and_xlsx_pictures(folder):
    pic = _png(folder / "p.png")
    assert len(list_embedded_images(_pptx(folder / "s.pptx", pic), "pptx")) == 1
    assert len(list_embedded_images(_xlsx(folder / "b.xlsx", pic), "xlsx")) == 1


def test_opendocument_and_epub_pictures(folder):
    png_bytes = (folder / "p.png").write_bytes(b"") or None
    buf = io.BytesIO()
    Image.new("RGB", (30, 30), "blue").save(buf, "PNG")
    data = buf.getvalue()
    odt = folder / "n.odt"
    with zipfile.ZipFile(odt, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", "<x/>")
        z.writestr("Pictures/100000000000.png", data)
    epub = folder / "b.epub"
    with zipfile.ZipFile(epub, "w") as z:
        z.writestr("mimetype", "application/epub+zip")
        z.writestr("META-INF/container.xml", "<c/>")
        z.writestr("OEBPS/images/cover.png", data)
        z.writestr("META-INF/ignored.png", data)            # not book content
    assert [im.label for im in list_embedded_images(odt, "odt")] == ["100000000000.png"]
    assert [im.label for im in list_embedded_images(epub, "epub")] == ["cover.png"]


def test_image_only_pdf_pages(folder):
    pdf = _scan_pdf(folder / "scan.pdf", pages=3)
    found = list_embedded_images(pdf, "pdf")
    assert [im.page for im in found] == [1, 2, 3]
    assert encode_for_model(found[0].load()).size == (400, 300)


def test_listing_never_raises(folder):
    junk = folder / "broken.docx"
    junk.write_bytes(b"PK\x03\x04 not really a zip")
    assert list_embedded_images(junk, "docx") == []
    assert list_embedded_images(folder / "nope.pdf", "pdf") == []


def test_page_ranges():
    assert page_ranges([1, 2, 3, 5, 7, 8]) == "1-3, 5, 7-8"
    assert page_ranges([4]) == "4"
    assert page_ranges([3, 1, 2, 2]) == "1-3"


def test_describe_embedded_notes(folder):
    # Distinct pictures: Word (python-docx) stores identical images only once.
    doc = _docx(folder / "d.docx", [_png(folder / "a.png"), _png(folder / "b.png", color=(0, 0, 200))])
    note = describe_embedded(list_embedded_images(doc, "docx"), doc, "docx")
    assert "2 embedded image(s)" in note and "image1.png" in note and "view_image" in note
    assert "index=1" in note and str(doc) in note
    pdf = _scan_pdf(folder / "s.pdf", pages=2)
    pnote = describe_embedded(list_embedded_images(pdf, "pdf"), pdf, "pdf")
    assert "page(s) 1-2" in pnote and "page=N" in pnote
    assert describe_embedded([], doc, "docx") == ""


# --- read_document now says when there are pictures --------------------------

def test_read_document_flags_pictures_at_the_top(folder):
    doc = _docx(folder / "Dashboard Tags.docx", [_png(folder / "a.png")])
    out = doc_tools.read_document({"path": str(doc)})
    header, _, body = out.partition("\n\n")
    assert "[Images:" in header and "1 embedded image(s)" in header
    assert "Dashboard tags for Canary" in out


def test_read_document_picture_only_docx_still_points_at_the_pictures(folder):
    doc = _docx(folder / "only.docx", [_png(folder / "a.png")], text="")
    out = doc_tools.read_document({"path": str(doc)})
    assert "(no extractable text)" in out and "view_image" in out


def test_read_document_scanned_pdf_offers_view_image(folder):
    pdf = _scan_pdf(folder / "scan.pdf", pages=2)
    out = doc_tools.read_document({"path": str(pdf)})
    assert "no text layer" in out and "its pages are images" in out
    assert "view_image" in out and "page=1" in out and not out.startswith("Error")


def test_read_document_without_pictures_has_no_note(folder):
    import docx
    d = docx.Document()
    d.add_paragraph("plain words")
    d.save(str(folder / "plain.docx"))
    assert "[Images:" not in doc_tools.read_document({"path": str(folder / "plain.docx")})


def test_read_file_redirects_images_to_view_image(folder):
    out = file_tools.read_file({"path": str(_png(folder / "a.png"))})
    assert "view_image" in out and out.startswith("Error")


# --- the view_image tool ------------------------------------------------------

def test_view_image_on_an_image_file(folder):
    out = view_image({"path": str(_png(folder / "a.png"))})
    assert isinstance(out, ToolOutput) and isinstance(out, str)
    assert len(out.images) == 1 and "a.png (200x100)" in out


def test_view_image_recognises_a_misnamed_image(folder):
    p = folder / "screenshot.dat"
    Image.new("RGB", (40, 40)).save(p, "PNG")
    assert len(view_image({"path": str(p)}).images) == 1


def test_view_image_on_a_docx(folder):
    doc = _docx(folder / "d.docx", [_png(folder / "a.png"), _png(folder / "b.png", size=(300, 100))])
    out = view_image({"path": str(doc)})
    assert len(out.images) == 2 and "(2 in total)" in out
    one = view_image({"path": str(doc), "index": 2})
    assert len(one.images) == 1 and "[2] image2.png 300x100" in one


def test_view_image_index_out_of_range(folder):
    doc = _docx(folder / "d.docx", [_png(folder / "a.png")])
    out = view_image({"path": str(doc), "index": 5})
    assert out.startswith("Error") and "1-1" in out and not getattr(out, "images", None)


def test_view_image_budget_one_tiled_picture_per_call(folder):
    big = _png(folder / "big.png", size=(1920, 1080))
    doc = _docx(folder / "d.docx", [big, _png(folder / "small.png")])
    out = view_image({"path": str(doc)})
    assert len(out.images) == 5                       # the tiled picture fills the call
    assert "index=2" in out                           # and says how to get the next


def test_identical_pictures_are_stored_once(folder):
    """Worth knowing: Word dedupes identical images, so 'image count' means distinct pictures."""
    doc = _docx(folder / "d.docx", [_png(folder / "a.png"), _png(folder / "b.png")])
    assert len(list_embedded_images(doc, "docx")) == 1


def test_view_image_small_pictures_share_a_call(folder):
    doc = _docx(folder / "d.docx", [_png(folder / f"{c}.png", color=(i * 60, 20, 20)) for i, c in enumerate("abc")])
    out = view_image({"path": str(doc)})
    assert len(out.images) == 3 and "More images" not in out


def test_view_image_pdf_pages(folder):
    pdf = _scan_pdf(folder / "scan.pdf", pages=3)
    out = view_image({"path": str(pdf), "page": 2})
    assert len(out.images) == 1 and "page 2" in out
    assert "no images" in view_image({"path": str(pdf), "page": 9})


def test_view_image_errors(folder):
    assert "File not found" in view_image({"path": str(folder / "nope.png")})
    assert "directory" in view_image({"path": str(folder)})
    assert "'path' is required" in view_image({})
    txt = folder / "notes.txt"
    txt.write_text("hello")
    assert "not an image" in view_image({"path": str(txt)})
    import docx
    d = docx.Document()
    d.save(str(folder / "plain.docx"))
    assert "no embedded images" in view_image({"path": str(folder / "plain.docx")})


def test_view_image_is_registered():
    from gemma_cli.tools import OLLAMA_TOOLS, execute_tool
    assert "view_image" in [t["function"]["name"] for t in OLLAMA_TOOLS]
    assert "'path' is required" in execute_tool("view_image", {})


def test_execute_tool_keeps_the_images(folder):
    from gemma_cli.tools import execute_tool
    out = execute_tool("view_image", {"path": str(_png(folder / "a.png"))})
    assert len(out.images) == 1


# --- how images reach Ollama ------------------------------------------------

class _Resp:
    def __init__(self, chunks):
        self._chunks = chunks

    def raise_for_status(self):
        pass

    def iter_lines(self):
        for c in self._chunks:
            yield json.dumps(c).encode()

    def close(self):
        pass


def _text(t):
    return [{"message": {"content": t}}, {"done": True}]


def _call(name, args):
    return [{"message": {"content": "", "tool_calls": [{"function": {"name": name, "arguments": args}}]}},
            {"done": True}]


def _fake_ollama(monkeypatch, script):
    sent = []

    def post(url, json=None, stream=False, timeout=None):
        sent.append(copy.deepcopy(json))
        return _Resp(script[min(len(sent) - 1, len(script) - 1)])

    monkeypatch.setattr(agent.requests, "post", post)
    return sent


BASE = {"ollama_url": "http://x", "model": "m", "num_ctx": 4096}


def test_view_image_pictures_ride_on_the_tool_message(folder, monkeypatch):
    png = _png(folder / "a.png")
    sent = _fake_ollama(monkeypatch, [_call("view_image", {"path": str(png)}), _text("FT-7314")])
    messages = [{"role": "system", "content": "s"}]
    events = list(agent.run_turn(dict(BASE), messages, "which tags are boxed?"))

    tool_msgs = [m for m in sent[1]["messages"] if m["role"] == "tool"]
    assert len(tool_msgs) == 1 and len(tool_msgs[0]["images"]) == 1
    assert type(tool_msgs[0]["content"]) is str                    # text, not the ToolOutput object
    results = [p for k, p in events if k == "tool_result"]
    assert results[0]["images"] == 1 and type(results[0]["result"]) is str
    # No fake user turn was injected to carry the picture.
    assert [m["role"] for m in messages] == ["system", "user", "assistant", "tool", "assistant"]


def test_user_attached_screenshot_is_tiled_and_explained(folder, monkeypatch):
    shot = _png(folder / "shot.png", size=(1920, 1080))
    sent = _fake_ollama(monkeypatch, [_text("ok")])
    list(agent.run_turn(dict(BASE), [{"role": "system", "content": "s"}], "read the tags", image_paths=[str(shot)]))
    user = sent[0]["messages"][-1]
    assert len(user["images"]) == 5
    assert user["content"].startswith("read the tags") and "zoomed tiles" in user["content"]


def test_missing_attachment_is_explained_not_silently_dropped(folder, monkeypatch):
    sent = _fake_ollama(monkeypatch, [_text("ok")])
    list(agent.run_turn(dict(BASE), [{"role": "system", "content": "s"}], "what is this?",
                        image_paths=[str(folder / "gone.png")]))
    user = sent[0]["messages"][-1]
    assert "images" not in user and "was not found" in user["content"]


# --- sessions and /check ------------------------------------------------------

def test_saved_sessions_drop_images_but_memory_keeps_them(folder):
    messages = [{"role": "system", "content": "s"},
                {"role": "user", "content": "look", "images": ["AAAA"]},
                {"role": "tool", "tool_name": "view_image", "content": "Showing a.png", "images": ["BBBB", "CCCC"]}]
    path = folder / "s.json"
    sessions.save_session(path, messages, "m")
    saved = json.loads(path.read_text(encoding="utf-8"))["messages"]
    assert all("images" not in m for m in saved)
    assert "1 image(s) were attached here" in saved[0]["content"]
    assert saved[1]["content"].startswith("Showing a.png") and "2 image(s)" in saved[1]["content"]
    assert messages[1]["images"] == ["AAAA"]                        # the live session is untouched


def test_check_sees_the_images_of_the_last_turn_only(monkeypatch):
    messages = [
        {"role": "system", "content": "s"},
        {"role": "user", "content": "old", "images": ["OLD"]},
        {"role": "assistant", "content": "old answer"},
        {"role": "user", "content": "which tags are boxed?"},
        {"role": "assistant", "content": "", "tool_calls": [{"function": {"name": "view_image", "arguments": {}}}]},
        {"role": "tool", "tool_name": "view_image", "content": "Showing a.png", "images": ["NEW1", "NEW2"]},
        {"role": "assistant", "content": "FT-7314"},
    ]
    assert review.last_turn_images(messages) == ["NEW1", "NEW2"]
    seen = {}

    def fake_chat_once(cfg, msgs, model, on_token=None):
        seen["msgs"] = msgs
        return "VERDICT: supported\nISSUES:\n- none"

    monkeypatch.setattr(agent, "_chat_once", fake_chat_once)
    review.check_last_turn(dict(BASE), messages, Console(file=io.StringIO(), no_color=True))
    request = seen["msgs"][0]
    assert request["images"] == ["NEW1", "NEW2"] and "part of the EVIDENCE" in request["content"]


def test_check_without_images_sends_none(monkeypatch):
    messages = [{"role": "system", "content": "s"}, {"role": "user", "content": "q"},
                {"role": "assistant", "content": "a"}]
    seen = {}
    monkeypatch.setattr(agent, "_chat_once", lambda cfg, msgs, model, on_token=None: seen.update(m=msgs) or "VERDICT: supported")
    review.check_last_turn(dict(BASE), messages, Console(file=io.StringIO(), no_color=True))
    assert "images" not in seen["m"][0]


# --- /image with spaces in the path -------------------------------------------

def _image_cmd(line):
    return main_mod._handle_command(line, {"model": "m"}, [{"role": "system", "content": ""}],
                                    Console(no_color=True, file=io.StringIO()), None)


def test_image_command_quoted_path_with_spaces(folder):
    sub = folder / "My Pictures"
    sub.mkdir()
    png = _png(sub / "shot 1.png")
    action, prompt, images = _image_cmd(f'/image "{png}" read the sign')
    assert action == "run" and prompt == "read the sign" and images == [str(png)]
    action, prompt, images = _image_cmd(f"/image '{png}' hi")
    assert images == [str(png)] and prompt == "hi"


def test_image_command_plain_path(folder):
    png = _png(folder / "a.png")
    action, prompt, images = _image_cmd(f"/image {png} what is this")
    assert action == "run" and images == [str(png)] and prompt == "what is this"


def test_image_command_rejects_missing_prompt_or_file(folder):
    png = _png(folder / "a.png")
    assert _image_cmd(f'/image "{png}"')[0] == "handled"
    assert _image_cmd(f"/image {folder / 'nope.png'} hi")[0] == "handled"
