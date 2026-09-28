"""Tests for write_spreadsheet and the write_file / edit_file guards.

The incident: asked for "an Excel sheet of the tags", the model tried pandas
(not installed) and fell back to a CSV — and the obvious alternative, write_file
to a .xlsx, would have written text into a file Excel refuses to open.

Workbooks are validated two ways: read back with openpyxl, and — independently —
by parsing the OOXML package with only zipfile + XML, so a pass does not rest on
openpyxl agreeing with itself.
"""
import copy
import csv
import io
import json
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
import pytest

from gemma_cli import agent
from gemma_cli.tools import doc_tools, file_tools, sheet_tools
from gemma_cli.tools.sheet_tools import rows_from_any, rows_from_text, write_spreadsheet

TAGS = [["Tag"], ["TT-9387"], ["LIC-3580"], ["HS-1492"], ["FV-4828"]]


@pytest.fixture
def folder(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    file_tools.set_allowed_write_roots([tmp_path])
    return tmp_path


def _values(path, sheet=None):
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet] if sheet else wb.active
    return [[c.value for c in row] for row in ws.iter_rows()]


def _ooxml_strings(path):
    """Every cell string in the first sheet, read with the standard library only."""
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        types = z.read("[Content_Types].xml").decode("utf-8")
        assert "spreadsheetml.sheet.main+xml" in types          # what Excel checks first
        assert "xl/workbook.xml" in names
        shared = []
        if "xl/sharedStrings.xml" in names:
            root = ET.fromstring(z.read("xl/sharedStrings.xml"))
            shared = ["".join(t.text or "" for t in si.iter(f"{{{ns['m']}}}t")) for si in root.findall("m:si", ns)]
        sheet = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
    out = []
    for c in sheet.iter(f"{{{ns['m']}}}c"):
        kind = c.get("t")
        if kind == "s":
            out.append(shared[int(c.find("m:v", ns).text)])
        elif kind == "inlineStr":
            out.append("".join(t.text or "" for t in c.iter(f"{{{ns['m']}}}t")))
        elif c.find("m:v", ns) is not None:
            out.append(c.find("m:v", ns).text)
    return out


# --- the tool ---------------------------------------------------------------

def test_writes_a_real_workbook(folder):
    out = write_spreadsheet({"path": str(folder / "tags.xlsx"), "rows": TAGS, "sheet": "Tags"})
    assert out.startswith("Created Excel workbook") and "5 row(s) x 1 column(s)" in out
    assert _values(folder / "tags.xlsx", "Tags") == TAGS
    assert _ooxml_strings(folder / "tags.xlsx") == [r[0] for r in TAGS]   # independent of openpyxl


def test_header_is_bold_and_frozen_and_columns_sized(folder):
    write_spreadsheet({"path": str(folder / "t.xlsx"),
                       "rows": [["Tag", "Description"], ["TT-9387", "reactor inlet temperature transmitter"]]})
    ws = openpyxl.load_workbook(folder / "t.xlsx").active
    assert ws["A1"].font.bold and not ws["A2"].font.bold
    assert ws.freeze_panes == "A2"
    assert ws.column_dimensions["B"].width > ws.column_dimensions["A"].width


def test_numbers_become_numbers_but_leading_zeros_stay_text(folder):
    write_spreadsheet({"path": str(folder / "n.xlsx"),
                       "rows": [["a", "b", "c", "d", "e", "f"], ["4200", "3.5", "007", "-12", 310, "TT-9387"]]})
    assert _values(folder / "n.xlsx")[1] == [4200, 3.5, "007", -12, 310, "TT-9387"]


def test_formulas_are_formulas(folder):
    out = write_spreadsheet({"path": str(folder / "f.xlsx"),
                             "rows": [["Item", "USD"], ["units", 4200], ["ship", 310], ["Total", "=SUM(B2:B3)"]]})
    wb = openpyxl.load_workbook(folder / "f.xlsx")
    ws = wb.active
    assert ws["B4"].value == "=SUM(B2:B3)" and ws["B4"].data_type == "f"
    # No cached value is written, so recalculation on open is what makes the
    # total appear - openpyxl sets fullCalcOnLoad by default; pin that.
    assert wb.calculation.fullCalcOnLoad
    assert "calculated when the file is opened" in out


@pytest.mark.parametrize("rows", [
    json.dumps(TAGS),                                                   # JSON string
    [{"Tag": r[0]} for r in TAGS[1:]],                                  # records
    {"Tag": [r[0] for r in TAGS[1:]]},                                  # dict of columns
    "| Tag |\n|---|\n| TT-9387 |\n| LIC-3580 |\n| HS-1492 |\n| FV-4828 |",   # markdown table
    "Tag\nTT-9387\nLIC-3580\nHS-1492\nFV-4828",                         # one-column text
])
def test_accepts_the_shapes_small_models_send(folder, rows):
    write_spreadsheet({"path": str(folder / "s.xlsx"), "rows": rows})
    assert _values(folder / "s.xlsx") == TAGS


def test_schema_asks_for_text_rows():
    """Nested-array arguments were the suspected reason a live turn ended silently;
    the schema asks for plain text, which the parser turns into rows."""
    from gemma_cli.tools.definitions import TOOLS
    spec = next(t for t in TOOLS if t["name"] == "write_spreadsheet")
    assert spec["input_schema"]["properties"]["rows"]["type"] == "string"


def test_csv_text_rows_with_quoted_commas(folder):
    text = 'Tag,Description\nTank$Farm/TIC101/PID.PV#Value,"BOILER 2 FEED, CONTROLLER"\n'
    write_spreadsheet({"path": str(folder / "q.xlsx"), "rows": text})
    assert _values(folder / "q.xlsx") == [["Tag", "Description"],
                                          ["Tank$Farm/TIC101/PID.PV#Value", "BOILER 2 FEED, CONTROLLER"]]


def test_flat_list_is_one_column(folder):
    write_spreadsheet({"path": str(folder / "flat.xlsx"), "rows": [r[0] for r in TAGS]})
    assert _values(folder / "flat.xlsx") == TAGS


def test_csv_and_tsv_text(folder):
    assert rows_from_text("Tag,Value\nTT-9387,12.5\n") == [["Tag", "Value"], ["TT-9387", "12.5"]]
    assert rows_from_text("Tag\tValue\nTT-9387\t12.5") == [["Tag", "Value"], ["TT-9387", "12.5"]]
    assert rows_from_any(None) == [] and rows_from_text("   ") == []


def test_headers_param_is_prepended_once(folder):
    write_spreadsheet({"path": str(folder / "h.xlsx"), "headers": ["Tag"], "rows": [["TT-9387"]]})
    assert _values(folder / "h.xlsx") == [["Tag"], ["TT-9387"]]
    write_spreadsheet({"path": str(folder / "h2.xlsx"), "headers": ["Tag"], "rows": [["Tag"], ["TT-9387"]]})
    assert _values(folder / "h2.xlsx") == [["Tag"], ["TT-9387"]]


def test_csv_output_has_a_bom_and_proper_quoting(folder):
    out = write_spreadsheet({"path": str(folder / "t.csv"), "rows": [["Tag", "Note"], ["TT-9387", "inlet, north"]]})
    assert "CSV file" in out
    raw = (folder / "t.csv").read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")                    # how Excel knows it is UTF-8
    assert list(csv.reader(io.StringIO(raw.decode("utf-8-sig")))) == [["Tag", "Note"], ["TT-9387", "inlet, north"]]


def test_missing_extension_defaults_to_xlsx(folder):
    write_spreadsheet({"path": str(folder / "tags"), "rows": TAGS})
    assert (folder / "tags.xlsx").exists()


def test_rejects_other_formats(folder):
    out = write_spreadsheet({"path": str(folder / "tags.docx"), "rows": TAGS})
    assert out.startswith("Error") and ".xlsx" in out and not (folder / "tags.docx").exists()


def test_existing_file_needs_overwrite_and_is_backed_up(folder):
    target = folder / "tags.xlsx"
    write_spreadsheet({"path": str(target), "rows": [["old"]]})
    refused = write_spreadsheet({"path": str(target), "rows": TAGS})
    assert refused.startswith("Error") and "overwrite=true" in refused
    assert _values(target) == [["old"]]                        # untouched
    out = write_spreadsheet({"path": str(target), "rows": TAGS, "overwrite": True})
    assert "backed up" in out and _values(target) == TAGS
    backups = list((folder / ".gemma" / "backups").glob("tags.xlsx.*.bak"))
    # openpyxl refuses a path ending in .bak, so open the backup as a stream.
    assert len(backups) == 1 and _values(io.BytesIO(backups[0].read_bytes())) == [["old"]]


def test_respects_write_roots(folder):
    allowed = folder / "allowed"
    allowed.mkdir()
    file_tools.set_allowed_write_roots([allowed])
    out = write_spreadsheet({"path": str(folder / "elsewhere.xlsx"), "rows": TAGS})
    assert "not permitted" in out and not (folder / "elsewhere.xlsx").exists()


def test_empty_rows_and_missing_path(folder):
    assert "empty" in write_spreadsheet({"path": str(folder / "e.xlsx"), "rows": []})
    assert "'path' is required" in write_spreadsheet({"rows": TAGS})


def test_sheet_names_are_made_valid(folder):
    write_spreadsheet({"path": str(folder / "s.xlsx"), "rows": TAGS,
                       "sheet": "Tags: [Q3]/2026 marked in red on the dashboard"})
    title = openpyxl.load_workbook(folder / "s.xlsx").active.title
    assert len(title) <= 31 and not any(ch in title for ch in "[]:*?/\\")


def test_file_open_in_excel_gets_a_useful_message(folder, monkeypatch):
    def locked(*a, **k):
        raise PermissionError(13, "Permission denied")
    monkeypatch.setattr(sheet_tools, "save_rows", locked)
    assert "open in Excel" in write_spreadsheet({"path": str(folder / "t.xlsx"), "rows": TAGS})


def test_read_document_reads_it_back(folder):
    write_spreadsheet({"path": str(folder / "tags.xlsx"), "rows": TAGS})
    out = doc_tools.read_document({"path": str(folder / "tags.xlsx")})
    assert all(r[0] in out for r in TAGS)


# --- write_file / edit_file guards ------------------------------------------------

def test_write_file_to_xlsx_makes_a_real_workbook(folder):
    out = file_tools.write_file({"path": str(folder / "tags.xlsx"), "content": "Tag\nTT-9387\nLIC-3580\n"})
    assert out.startswith("Wrote a real Excel workbook") and "write_spreadsheet" in out
    assert (folder / "tags.xlsx").read_bytes()[:2] == b"PK"    # a zip package, not text
    assert _values(folder / "tags.xlsx") == [["Tag"], ["TT-9387"], ["LIC-3580"]]
    assert _ooxml_strings(folder / "tags.xlsx") == ["Tag", "TT-9387", "LIC-3580"]


def test_write_file_to_xlsx_accepts_a_list(folder):
    file_tools.write_file({"path": str(folder / "t.xlsx"), "content": TAGS})
    assert _values(folder / "t.xlsx") == TAGS


@pytest.mark.parametrize("name", ["report.docx", "deck.pptx", "out.pdf", "old.xls", "m.xlsm", "n.odt"])
def test_write_file_refuses_binary_office_formats(folder, name):
    out = file_tools.write_file({"path": str(folder / name), "content": "hello"})
    assert out.startswith("Error") and "plain text" in out and not (folder / name).exists()


def test_write_file_refusal_points_spreadsheets_at_write_spreadsheet(folder):
    assert "write_spreadsheet" in file_tools.write_file({"path": str(folder / "old.xls"), "content": "x"})
    assert ".md" in file_tools.write_file({"path": str(folder / "r.docx"), "content": "x"})


def test_write_file_non_string_content_to_text_file(folder):
    file_tools.write_file({"path": str(folder / "d.json"), "content": {"a": 1}})
    assert json.loads((folder / "d.json").read_text()) == {"a": 1}


def test_edit_file_on_a_workbook_explains_instead_of_a_codec_error(folder):
    write_spreadsheet({"path": str(folder / "tags.xlsx"), "rows": TAGS})
    out = file_tools.edit_file({"path": str(folder / "tags.xlsx"), "old_string": "TT", "new_string": "XX"})
    assert out.startswith("Error") and "write_spreadsheet" in out and "codec" not in out


# --- registration, approval gate, read-only children ------------------------------

def test_registered_and_dispatched(folder):
    from gemma_cli.tools import OLLAMA_TOOLS, execute_tool
    assert "write_spreadsheet" in [t["function"]["name"] for t in OLLAMA_TOOLS]
    assert "Created Excel workbook" in execute_tool("write_spreadsheet", {"path": str(folder / "x.xlsx"), "rows": TAGS})


def test_one_mutating_list_covers_approval_and_children():
    assert "write_spreadsheet" in agent.MUTATING_TOOLS


class _Resp:
    def __init__(self, chunks):
        self._chunks = chunks

    def raise_for_status(self):
        pass

    def iter_lines(self):
        for c in self._chunks:
            yield json.dumps(c).encode()

    def close(self):
        pass


def _fake(monkeypatch, script):
    sent = []

    def post(url, json=None, stream=False, timeout=None):
        sent.append(copy.deepcopy(json))
        return _Resp(script[min(len(sent) - 1, len(script) - 1)])

    monkeypatch.setattr(agent.requests, "post", post)
    return sent


def test_approval_mode_gates_write_spreadsheet(folder, monkeypatch):
    call = [{"message": {"content": "", "tool_calls": [{"function": {"name": "write_spreadsheet",
             "arguments": {"path": str(folder / "t.xlsx"), "rows": TAGS}}}]}}, {"done": True}]
    _fake(monkeypatch, [call, [{"message": {"content": "ok"}}, {"done": True}]])
    asked = []
    list(agent.run_turn({"ollama_url": "http://x", "model": "m"}, [{"role": "system", "content": "s"}], "go",
                        approver=lambda name, args: asked.append(name) or False))
    assert asked == ["write_spreadsheet"] and not (folder / "t.xlsx").exists()


def test_read_only_children_cannot_write_spreadsheets(folder, monkeypatch):
    sent = _fake(monkeypatch, [[{"message": {"content": "ok"}}, {"done": True}]])
    list(agent.run_child({"ollama_url": "http://x", "model": "m"}, "x", readonly=True))
    assert "write_spreadsheet" not in {t["function"]["name"] for t in sent[0]["tools"]}
