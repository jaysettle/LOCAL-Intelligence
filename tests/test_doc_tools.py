"""Tests for read_document.

Every fixture builds a REAL file of the format under test (a real zip container,
a real PDF byte stream, a real docx via python-docx) and runs it through the
public tool, so these catch backend API drift rather than asserting on mocks.
No network and no Ollama required.
"""
import json
import zipfile

import pytest

from gemma_cli.tools import doc_tools, file_tools
from gemma_cli.tools.doc_tools import read_document


# --- fixture builders ----------------------------------------------------

def _make_pdf(path, pages):
    """A minimal but valid multi-page PDF with a text layer, built by hand.

    Avoids a test-only dependency on a PDF writer; pypdf must be able to pull
    the text back out of the content streams.
    """
    objects = {}
    kids = []
    obj_num = 10  # page/content objects start above the fixed 1,2,5 objects
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
        content_num = obj_num + 1
        objects[obj_num] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 5 0 R >> >> /Contents {content_num} 0 R >>"
        ).encode("latin-1")
        objects[content_num] = (
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
        kids.append(f"{obj_num} 0 R")
        obj_num += 2

    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[2] = (
        f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>"
    ).encode("latin-1")
    objects[5] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for num in sorted(objects):
        offsets[num] = len(out)
        out += f"{num} 0 obj\n".encode() + objects[num] + b"\nendobj\n"
    xref_at = len(out)
    highest = max(objects) + 1
    out += f"xref\n0 {highest}\n".encode()
    out += b"0000000000 65535 f \n"
    for num in range(1, highest):
        out += (f"{offsets[num]:010d} 00000 n \n").encode() if num in offsets else b"0000000000 65535 f \n"
    out += f"trailer\n<< /Size {highest} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))
    return path


def _make_odf(path, mimetype, content_xml):
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", mimetype)
        z.writestr("META-INF/manifest.xml", "<manifest/>")
        z.writestr("content.xml", content_xml)
    return path


_ODF_NS = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"'
)


@pytest.fixture
def docs(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    file_tools.set_allowed_write_roots([tmp_path])
    return tmp_path


# --- PDF -----------------------------------------------------------------

def test_pdf_text_and_page_labels(docs):
    pdf = _make_pdf(docs / "report.pdf", ["Quarterly revenue rose", "Costs fell sharply"])
    out = read_document({"path": str(pdf)})
    assert "[pdf, 2 pages]" in out
    assert "--- page 1 ---" in out and "--- page 2 ---" in out
    assert "Quarterly revenue rose" in out
    assert "Costs fell sharply" in out


def test_pdf_offset_and_limit(docs):
    pdf = _make_pdf(docs / "many.pdf", [f"Page number {i}" for i in range(1, 6)])
    out = read_document({"path": str(pdf), "offset": 3, "limit": 2})
    assert "Page number 3" in out and "Page number 4" in out
    assert "Page number 1" not in out and "Page number 5" not in out
    assert "showing pages 3-4 of 5" in out


def test_pdf_offset_past_end_explains(docs):
    pdf = _make_pdf(docs / "short.pdf", ["only page"])
    out = read_document({"path": str(pdf), "offset": 9})
    assert "past the end" in out and "between 1 and 1" in out


def test_pdf_without_text_layer_says_so(docs):
    pdf = _make_pdf(docs / "scan.pdf", [""])
    out = read_document({"path": str(pdf)})
    assert "no text layer" in out.lower()
    assert "ocr" in out.lower()


# --- Word ----------------------------------------------------------------

def test_docx_paragraphs_headings_and_tables(docs):
    import docx

    d = docx.Document()
    d.add_heading("Project Status", level=1)
    d.add_paragraph("The migration finished on time.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Total"
    table.cell(1, 0).text = "North"
    table.cell(1, 1).text = "412"
    path = docs / "status.docx"
    d.save(str(path))

    out = read_document({"path": str(path)})
    assert "[docx" in out
    assert "# Project Status" in out
    assert "The migration finished on time." in out
    assert "Region | Total" in out
    assert "North | 412" in out


# --- Excel ---------------------------------------------------------------

def test_xlsx_multiple_sheets(docs):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Item", "Qty"])
    ws.append(["Widget", 7])
    wb.create_sheet("Notes").append(["reviewed by finance"])
    path = docs / "book.xlsx"
    wb.save(str(path))

    out = read_document({"path": str(path)})
    assert "[xlsx, 2 sheets]" in out
    assert "--- sheet 'Sales' ---" in out
    assert "Widget | 7" in out
    assert "reviewed by finance" in out


def test_xlsx_row_cap_reports_remainder(docs):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(300):
        ws.append([f"row{i}"])
    path = docs / "big.xlsx"
    wb.save(str(path))

    out = read_document({"path": str(path), "max_chars": 100000})
    assert "more rows not shown" in out


# --- PowerPoint ----------------------------------------------------------

def test_pptx_slides_and_notes(docs):
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Roadmap"
    slide.placeholders[1].text = "Ship the parser"
    slide.notes_slide.notes_text_frame.text = "mention the timeline"
    path = docs / "deck.pptx"
    prs.save(str(path))

    out = read_document({"path": str(path)})
    assert "[pptx, 1 slides]" in out
    assert "Roadmap" in out
    assert "Ship the parser" in out
    assert "[speaker notes] mention the timeline" in out


# --- OpenDocument --------------------------------------------------------

def test_odt_text(docs):
    content = (
        f'<office:document-content {_ODF_NS}><office:body><office:text>'
        "<text:h>Chapter One</text:h>"
        "<text:p>It was a bright cold day.</text:p>"
        "</office:text></office:body></office:document-content>"
    )
    path = _make_odf(docs / "novel.odt", "application/vnd.oasis.opendocument.text", content)
    out = read_document({"path": str(path)})
    assert "[odt" in out
    assert "Chapter One" in out
    assert "It was a bright cold day." in out


def test_ods_rows(docs):
    content = (
        f'<office:document-content {_ODF_NS}><office:body><office:spreadsheet>'
        "<table:table><table:table-row>"
        "<table:table-cell><text:p>alpha</text:p></table:table-cell>"
        "<table:table-cell><text:p>beta</text:p></table:table-cell>"
        "</table:table-row></table:table>"
        "</office:spreadsheet></office:body></office:document-content>"
    )
    path = _make_odf(docs / "sheet.ods", "application/vnd.oasis.opendocument.spreadsheet", content)
    out = read_document({"path": str(path)})
    assert "alpha" in out and "beta" in out


# --- RTF / EPUB / email / notebook / CSV / HTML --------------------------

def test_rtf(docs):
    path = docs / "memo.rtf"
    path.write_text(r"{\rtf1\ansi Hello from \b RTF\b0 land.}", encoding="utf-8")
    out = read_document({"path": str(path)})
    assert "[rtf" in out
    assert "Hello from" in out and "RTF" in out


def test_epub_spine_order(docs):
    path = docs / "book.epub"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/epub+zip")
        z.writestr(
            "META-INF/container.xml",
            '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
            '<rootfiles><rootfile full-path="OEBPS/book.opf"/></rootfiles></container>',
        )
        z.writestr(
            "OEBPS/book.opf",
            '<package xmlns="http://www.idpf.org/2007/opf"><manifest>'
            '<item id="c1" href="one.xhtml"/><item id="c2" href="two.xhtml"/>'
            "</manifest><spine><itemref idref=\"c1\"/><itemref idref=\"c2\"/></spine></package>",
        )
        z.writestr("OEBPS/one.xhtml", "<html><body><p>First chapter body</p></body></html>")
        z.writestr("OEBPS/two.xhtml", "<html><body><p>Second chapter body</p></body></html>")

    out = read_document({"path": str(path)})
    assert "[epub, 2 chapters]" in out
    assert out.index("First chapter body") < out.index("Second chapter body")


def test_eml_headers_and_body(docs):
    path = docs / "note.eml"
    path.write_text(
        "From: a@example.com\nTo: b@example.com\nSubject: Budget\n"
        "Content-Type: text/plain\n\nThe numbers are attached.\n",
        encoding="utf-8",
    )
    out = read_document({"path": str(path)})
    assert "Subject: Budget" in out
    assert "The numbers are attached." in out


def test_ipynb_cells_and_outputs(docs):
    nb = {
        "cells": [
            {"cell_type": "markdown", "source": ["# Analysis\n"]},
            {
                "cell_type": "code",
                "source": ["print('hi')\n"],
                "outputs": [{"text": ["hi\n"]}],
            },
        ]
    }
    path = docs / "nb.ipynb"
    path.write_text(json.dumps(nb), encoding="utf-8")
    out = read_document({"path": str(path)})
    assert "# Analysis" in out
    assert "print('hi')" in out
    assert "[output 2] hi" in out


def test_csv_rows_and_count(docs):
    path = docs / "data.csv"
    path.write_text("name,score\nada,99\ngrace,98\n", encoding="utf-8")
    out = read_document({"path": str(path)})
    assert "name | score" in out
    assert "ada | 99" in out
    assert "3 rows" in out


def test_html_strips_markup(docs):
    path = docs / "page.html"
    path.write_text(
        "<html><head><style>p{color:red}</style></head>"
        "<body><h1>Title here</h1><p>Body text here</p></body></html>",
        encoding="utf-8",
    )
    out = read_document({"path": str(path)})
    assert "Title here" in out and "Body text here" in out
    assert "color:red" not in out


# --- detection, limits, errors -------------------------------------------

def test_content_beats_extension(docs):
    """A .docx renamed to .doc must still be read as a .docx, with a note."""
    import docx

    d = docx.Document()
    d.add_paragraph("mislabelled but readable")
    path = docs / "renamed.doc"
    d.save(str(path))

    out = read_document({"path": str(path)})
    assert "mislabelled but readable" in out
    assert "the extension says .doc but the content is docx" in out


def test_legacy_ole_without_libreoffice_explains(docs, monkeypatch):
    monkeypatch.setattr(doc_tools, "_find_soffice", lambda: None)
    path = docs / "old.doc"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    out = read_document({"path": str(path)})
    assert "legacy Office" in out
    assert "Save As" in out and "LibreOffice" in out


def test_max_chars_truncates_and_points_forward(docs):
    pdf = _make_pdf(docs / "long.pdf", [f"Page body {i} " + "x" * 400 for i in range(1, 8)])
    out = read_document({"path": str(pdf), "max_chars": 900})
    assert len(out) < 1600
    assert "call read_document again with offset=" in out


def test_missing_file(docs):
    out = read_document({"path": str(docs / "nope.pdf")})
    assert "File not found" in out


def test_directory_rejected(docs):
    out = read_document({"path": str(docs)})
    assert "is a directory" in out


def test_missing_path_arg():
    assert "'path' is required" in read_document({})


def test_unknown_binary_format(docs):
    path = docs / "mystery.bin"
    path.write_bytes(b"\x07\x08\x09not a document")
    out = read_document({"path": str(path)})
    # Falls through to the text reader rather than erroring — but a truly
    # unknown *extension* with no sniffable content is still handled sanely.
    assert "mystery.bin" in out


def test_corrupt_docx_reports_cleanly(docs):
    """A zip that claims to be a docx but is not must not raise."""
    path = docs / "broken.docx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml", "<not-valid-ooxml>")
    out = read_document({"path": str(path)})
    assert out.startswith("Error reading broken.docx")


# --- read_file redirect ---------------------------------------------------

def test_read_file_redirects_binary_documents(docs):
    pdf = _make_pdf(docs / "redirect.pdf", ["hello"])
    out = file_tools.read_file({"path": str(pdf)})
    assert "read_document" in out
    assert "cannot decode" in out


def test_read_file_still_reads_plain_text(docs):
    path = docs / "notes.txt"
    path.write_text("plain content\n", encoding="utf-8")
    assert "plain content" in file_tools.read_file({"path": str(path)})


# --- tool registration ----------------------------------------------------

def test_tool_is_registered_and_dispatchable():
    from gemma_cli.tools import OLLAMA_TOOLS, execute_tool

    names = [t["function"]["name"] for t in OLLAMA_TOOLS]
    assert "read_document" in names
    assert "'path' is required" in execute_tool("read_document", {})
